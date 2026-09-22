"""
Prep helper (not part of the numbered pipeline): converts Ipsos's MRP
"Constituency data" sheet into the CSV shape 04_ingest_mrp_release.py
expects. Nicest format seen so far — one row per seat, decimal-fraction
vote shares per party, AND its own pcon_code column, so no fuzzy name
matching should be needed... except 5 Scottish seats (Ayr Carrick and
Cumnock, Berwickshire Roxburgh and Selkirk, Central Ayrshire, Kilmarnock
and Loudoun, West Aberdeenshire and Kincardine) use different pcon_codes
than our `constituencies` table has for the identically-named seats —
likely a pre-final ONS numbering Ipsos had at the time. Rather than crash
on the foreign-key constraint (schema.sql has pcon_code NOT NULL
REFERENCES constituencies), this checks each code against the DB and
falls back to passing the name instead (blank pcon_code) for
04_ingest_mrp_release.py's own fuzzy matcher to resolve — which will be
an exact string match for these 5, since only the code differs.

Ignores the confidence-interval columns (_low/_high) and "Seat category"
— 04_ingest_mrp_release.py's schema only stores a point estimate per
party, not a range.

Usage:
    python scripts/prep_ipsos_xlsx.py --xlsx /path/to/ipsos-mrp-2024-constituency-estimates-data.xlsx --out data/mrp_prepped/ipsos_2024-06-18.csv
"""
import argparse
import csv
import sqlite3
import sys
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DB_PATH = ROOT / "data" / "uk_elections.db"

PARTY_MAP = {"con": "Con", "lab": "Lab", "ld": "LD", "ref": "RUK", "grn": "Green", "oth": "Other", "pc": "PC", "snp": "SNP"}


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--xlsx", required=True)
    parser.add_argument("--sheet", default="Constituency data")
    parser.add_argument("--out", required=True)
    args = parser.parse_args()

    con = sqlite3.connect(DB_PATH)
    known_codes = set(r[0] for r in con.execute("SELECT pcon_code FROM constituencies"))
    con.close()

    wb = openpyxl.load_workbook(args.xlsx, data_only=True)
    ws = wb[args.sheet]
    header = [str(c).strip() if c else "" for c in next(ws.iter_rows(min_row=1, max_row=1, values_only=True))]

    est_cols = {}  # our_code -> column index, only the "_est" (point estimate) columns
    for i, name in enumerate(header):
        key = name.strip().lower()
        if key.endswith("_est"):
            prefix = key[:-4]
            if prefix in PARTY_MAP:
                est_cols[PARTY_MAP[prefix]] = i

    code_col = header.index("pcon")
    name_col = header.index("pcon_name")

    out_rows = []
    n_seats = 0
    n_code_mismatch = 0
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, values_only=True):
        if not row[code_col]:
            continue
        n_seats += 1
        code, name = row[code_col], row[name_col]
        if code not in known_codes:
            n_code_mismatch += 1
            code = ""  # fall back to name matching in 04_ingest_mrp_release.py
        for our_party, i in est_cols.items():
            share = row[i]
            if isinstance(share, (int, float)) and share:
                out_rows.append({"pcon_code": code, "pcon_name": name, "party": our_party,
                                  "vote_share_pct": round(share * 100, 3), "win_probability_pct": ""})

    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    with out_path.open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=["pcon_code", "pcon_name", "party", "vote_share_pct", "win_probability_pct"])
        writer.writeheader()
        writer.writerows(out_rows)

    print(f"Wrote {len(out_rows)} party rows across {n_seats} seats to {out_path} "
          f"({n_code_mismatch} seats had a pcon_code not in our DB, falling back to name matching)")


if __name__ == "__main__":
    sys.exit(main())
